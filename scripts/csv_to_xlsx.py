"""Turn a diagnostics CSV into a formatted xlsx.

The CSVs in results/diagnostics are written for machines: full float precision,
bare column names, no units. This makes one readable in Excel - friendly
headers, sensible number formats, frozen header row, autosized columns - without
touching the CSV, which stays the machine-readable record.

Run:
    python scripts/csv_to_xlsx.py results/diagnostics/hvdc_effect.csv
    python scripts/csv_to_xlsx.py <csv> [<out.xlsx>]
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Column -> (display name, Excel number format). Anything unlisted keeps its
# name and is left to Excel's general format.
HEADERS = {
    "config": ("Configuration", None),
    "hvdc": ("HVDC read?", None),
    "no_dc_load": ("DC terminal load excluded?", None),
    "components": ("Connected components", "#,##0"),
    "largest_island": ("Largest island (buses)", "#,##0"),
    "singletons": ("Single-bus islands", "#,##0"),
    "total_load_MW": ("Total load (MW)", "#,##0"),
    "structural_floor_MW": ("Structural floor (MW)", "#,##0.0"),
    "EAENS_conn": ("EAENS, connectivity (MWh)", "#,##0.000"),
    "EAENS_dcopf": ("EAENS, DC OPF (MWh)", "#,##0.000"),
    "build_s": ("Model build (s)", "#,##0.0"),
    # state_summary.csv columns, so the same tool handles those too
    "stamp": ("Timestamp (UTC)", None),
    "label": ("Hour", None),
    "disagg": ("Load disaggregation", None),
    "lines_over_100": ("AC lines over 100%", "#,##0"),
    "max_loading_pct": ("Worst line (% of rating)", "#,##0"),
    "shed_MW": ("Load shed (MW)", "#,##0"),
    "shed_pct_of_demand": ("Shed (% of demand)", "0.00%"),
    "dc_saturated": ("HVDC elements at cap", "#,##0"),
    "thermal_limits": ("Thermal limits", None),
    "rating_factor": ("Rating factor", "0.0#"),
    "objective": ("Objective", None),
    "backend": ("Solver", None),
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")


def convert(src: Path, dst: Path) -> None:
    df = pd.read_csv(src)
    # Excel renders a true fraction as a percentage; the CSV stores 0-100.
    pct = [c for c in df.columns if HEADERS.get(c, ("", ""))[1] == "0.00%"]
    for c in pct:
        df[c] = df[c] / 100.0
    df = df.rename(columns={k: v[0] for k, v in HEADERS.items() if k in df.columns})

    with pd.ExcelWriter(dst, engine="openpyxl") as xl:
        df.to_excel(xl, index=False, sheet_name="results")
        ws = xl.sheets["results"]
        fmts = {HEADERS[k][0]: HEADERS[k][1] for k in HEADERS}
        for j, col in enumerate(df.columns, start=1):
            letter = get_column_letter(j)
            cell = ws.cell(row=1, column=j)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center")
            nf = fmts.get(col)
            if nf:
                for i in range(2, len(df) + 2):
                    ws.cell(row=i, column=j).number_format = nf
            body = max((len(str(v)) for v in df[col].astype(str)), default=0)
            ws.column_dimensions[letter].width = min(
                max(11, body + 2, min(len(col) + 2, 18)), 30)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 34
        ws.auto_filter.ref = ws.dimensions
    print(f"wrote {dst}  ({len(df)} rows x {len(df.columns)} columns)")


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        raise SystemExit(2)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".xlsx")
    if not src.is_file():
        raise SystemExit(f"no such file: {src}")
    convert(src, dst)


if __name__ == "__main__":
    main()
